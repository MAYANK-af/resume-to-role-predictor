"""
server.py
FastAPI backend wrapper for Resume-to-Role Predictor.
Keeps existing model code untouched. Exposes POST /predict and metadata endpoints.
"""

import os
import io
from typing import Optional, List, Dict
from fastapi import FastAPI, HTTPException, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

from src.predict import get_predictor

app = FastAPI(
    title="Resume-to-Role Predictor API",
    description="3D Word Constellation & Role Alignment Engine",
    version="2.0.0"
)

# Enable CORS for local Vite development and preview servers
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Pre-load model predictor
predictor = get_predictor()


class PredictRequest(BaseModel):
    resume_text: Optional[str] = Field(None, description="Raw text of the resume to analyze")
    text: Optional[str] = Field(None, description="Alias for resume_text")
    target_role: Optional[str] = Field(None, description="Optional target role to score against")

    def get_text(self) -> str:
        return self.resume_text or self.text or ""


class WordWeight(BaseModel):
    word: str
    weight: float
    tfidf: Optional[float] = None
    coef: Optional[float] = None


class PredictResponse(BaseModel):
    predicted_role: str
    nb_predicted_role: str
    target_role: str
    fit_score: float
    lr_prob_fit_score: float
    probabilities: Dict[str, float]
    top_words: List[WordWeight]
    target_words: List[WordWeight]
    cleaned_tokens: List[str]


class UploadResponse(BaseModel):
    filename: str
    text: str
    word_count: int
    char_count: int
    file_type: str



# Pre-packaged editorial sample profiles
SAMPLE_PROFILES = {
    "AI / ML": (
        "Senior Data Scientist and Machine Learning Engineer with 5+ years of experience building and deploying "
        "NLP and deep learning systems. Proficient in Python, pandas, numpy, scikit-learn, PyTorch, and TensorFlow. "
        "Developed transformer-based text classification pipelines, clustering algorithms, and recommendation engines. "
        "Expertise in exploratory data analysis, feature engineering, regression modeling, hyperparameter tuning, "
        "model explainability with SHAP, and SQL for querying relational databases. Designed automated ETL pipelines "
        "and monitored predictive ML models in production environments."
    ),
    "Backend": (
        "Senior Backend Developer with 6 years of experience architecting high-performance enterprise applications "
        "and distributed microservices. Core expertise in Java, Spring Boot, Hibernate, Python, and SQL Server. "
        "Designed and implemented RESTful APIs, optimized complex database queries in PostgreSQL and Oracle, "
        "and integrated Redis caching layers. Extensive experience in Docker, Kubernetes, CI/CD pipelines with Jenkins "
        "and GitLab, and messaging queues using Apache Kafka and RabbitMQ. Strong background in database design and OOP."
    ),
    "Web Dev": (
        "Creative Front-End Web Designer and Developer with 4 years of experience crafting modern, responsive web user interfaces. "
        "Proficient in HTML5, CSS3, JavaScript, jQuery, Bootstrap, React.js, and WordPress theme customization. "
        "Skilled in Adobe Photoshop, Figma, graphic design, responsive grid layouts, and cross-browser compatibility. "
        "Developed client websites, optimized UI/UX design workflows, and integrated REST APIs. Strong knowledge of web "
        "accessibility standards, animation effects, and mobile-first responsive design."
    ),
    "Analyst": (
        "Business Analyst with 5 years experience driving business intelligence, requirement gathering, and process re-engineering. "
        "Expert in translating complex business stakeholder requirements into functional specifications, user stories, and wireframes. "
        "Proficient in SQL, Excel advanced financial modeling, Tableau, Power BI, and ERP systems. "
        "Conducted GAP analysis, cost-benefit analysis, executive reporting, and KPI dashboard development. "
        "Facilitated sprint planning, client requirement workshops, and user acceptance testing."
    ),
    "Other Tech": (
        "QA Automation Engineer and Software Tester with 5 years experience in manual and automated testing of web applications. "
        "Proficient in Selenium WebDriver, TestNG, Java, Python, and Postman API testing. Built robust test automation frameworks, "
        "executed regression testing suites, performance testing, and bug tracking using JIRA. Experience in network security protocols, "
        "vulnerability scanning, and blockchain smart contract testing. Certified ISTQB software quality engineer."
    )
}


@app.get("/health")
def health_check():
    return {"status": "healthy", "service": "Resume-to-Role Predictor API"}


@app.get("/metadata")
def get_metadata():
    """Returns role taxonomy and global top feature coefficients for interactive 3D inspection."""
    meta = predictor.metadata
    return {
        "roles": predictor.roles,
        "top_coefficients": meta.get("top_coefficients", {}),
        "comparison_metrics": meta.get("comparison_df", {}).to_dict(orient="records") if hasattr(meta.get("comparison_df"), "to_dict") else meta.get("comparison_df"),
        "regression_comparison": meta.get("regression_comparison", {}).to_dict(orient="records") if hasattr(meta.get("regression_comparison"), "to_dict") else meta.get("regression_comparison")
    }


@app.get("/samples")
def get_samples():
    """Returns realistic curated technical resumes for zero-friction user testing."""
    return SAMPLE_PROFILES


@app.post("/predict", response_model=PredictResponse)
def predict_resume(request: PredictRequest):
    """
    Analyzes resume text:
    - Validates minimum word length
    - Classifies best-fit tech role
    - Computes 0-100 fit score (Ridge Regression)
    - Returns per-role probabilities and top influential word contributions
    """
    input_text = request.get_text()
    is_valid, err_msg = predictor.validate_input(input_text)
    if not is_valid:
        raise HTTPException(status_code=400, detail=err_msg)

    try:
        raw_res = predictor.predict(input_text, target_role=request.target_role)

        # Map top words to {word, weight} as required by prompt
        top_words = [
            WordWeight(
                word=w["display"],
                weight=w["contribution"],
                tfidf=w.get("tfidf"),
                coef=w.get("coef")
            )
            for w in raw_res["top_driving_words"]
        ]

        target_words = [
            WordWeight(
                word=w["display"],
                weight=w["contribution"],
                tfidf=w.get("tfidf"),
                coef=w.get("coef")
            )
            for w in raw_res["target_driving_words"]
        ]

        return PredictResponse(
            predicted_role=raw_res["predicted_role"],
            nb_predicted_role=raw_res["nb_predicted_role"],
            target_role=raw_res["target_role"],
            fit_score=raw_res["ridge_fit_score"],
            lr_prob_fit_score=raw_res["lr_prob_fit_score"],
            probabilities=raw_res["probabilities"],
            top_words=top_words,
            target_words=target_words,
            cleaned_tokens=raw_res["cleaned_tokens"]
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


def extract_text_from_file(filename: str, content: bytes) -> str:
    """
    Extracts plain text content from uploaded files (.pdf, .docx, .txt, .md).
    Handles encoding variations and clean text stitching.
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content))
            text_parts = []
            for page in reader.pages:
                txt = page.extract_text()
                if txt:
                    text_parts.append(txt)
            full_text = "\n\n".join(text_parts).strip()
            if not full_text:
                raise ValueError("No readable text found in PDF. Scanned image-only PDFs or encrypted PDFs cannot be parsed directly.")
            return full_text
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"Failed to read PDF document: {str(e)}")

    elif ext in [".docx", ".doc"]:
        if ext == ".docx":
            try:
                import docx
                doc = docx.Document(io.BytesIO(content))
                text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                text_parts.append(cell.text.strip())
                full_text = "\n\n".join(text_parts).strip()
                if not full_text:
                    raise ValueError("Word document (.docx) contains no readable text.")
                return full_text
            except ValueError:
                raise
            except Exception as e:
                raise ValueError(f"Failed to parse Word document: {str(e)}")
        else:
            raise ValueError("Legacy binary .doc format is not supported directly. Please save or export as .docx or .pdf.")

    elif ext in [".txt", ".md", ".rtf", ".text"]:
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                decoded = content.decode(encoding).strip()
                if decoded:
                    return decoded
            except UnicodeDecodeError:
                continue
        raise ValueError("Could not decode text file. Ensure it is encoded in UTF-8 or standard ASCII.")

    else:
        raise ValueError(f"Unsupported file format '{ext}'. Please upload a PDF (.pdf), Word document (.docx), or plain text file (.txt).")


@app.post("/upload-resume", response_model=UploadResponse)
@app.post("/upload", response_model=UploadResponse)
async def upload_resume_file(file: UploadFile = File(...)):
    """
    Extracts text from uploaded resume documents (.pdf, .docx, .txt, .md).
    Returns parsed plain text and token metadata for immediate analysis.
    """
    if not file.filename:
        raise HTTPException(status_code=400, detail="No file was provided.")

    content = await file.read()
    if len(content) == 0:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    # Limit file size to 10MB
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File size exceeds the 10MB maximum limit.")

    try:
        extracted_text = extract_text_from_file(file.filename, content)
        words = extracted_text.split()
        ext = os.path.splitext(file.filename)[1].lower().replace(".", "")
        return UploadResponse(
            filename=file.filename,
            text=extracted_text,
            word_count=len(words),
            char_count=len(extracted_text),
            file_type=ext.upper() if ext else "TEXT"
        )
    except ValueError as ve:
        raise HTTPException(status_code=400, detail=str(ve))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error extracting text from file: {str(e)}")


# Mount compiled 3D React frontend if built
dist_dir = os.path.join(os.path.dirname(__file__), "frontend", "dist")
if os.path.exists(dist_dir):
    from fastapi.staticfiles import StaticFiles
    app.mount("/", StaticFiles(directory=dist_dir, html=True), name="frontend")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("server:app", host="127.0.0.1", port=8000, reload=False)

